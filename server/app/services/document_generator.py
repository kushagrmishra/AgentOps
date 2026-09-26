from __future__ import annotations

import html
import json
import logging
import os
import re
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from app.core.config import settings

logger = logging.getLogger(__name__)


# =============================================================================
# 1. PARSED REPORT DATA STRUCTURE & HELPER
# =============================================================================

@dataclass
class ReportData:
    title: str
    subtitle: str
    period: str
    entity: str
    classification: str
    kpis: list[dict[str, str]]
    covenants: list[dict[str, str]]
    concentration: list[dict[str, str]]
    risks: list[dict[str, str]]
    recommendations: list[str]


def clean_inline_md(text: str) -> str:
    """Strip markdown links and formatting for plain text use in PPTX/DOCX runs."""
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = re.sub(r"[*_]{1,2}([^*_]+)[*_]{1,2}", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def parse_markdown_blocks(text: str) -> list[dict[str, Any]]:
    """Parse raw markdown into structured blocks (heading, paragraph, list, table, blockquote)."""
    blocks: list[dict[str, Any]] = []
    lines = text.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        # Heading
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading_title = stripped.lstrip("# ").strip()
            blocks.append({"type": "heading", "level": level, "text": heading_title})
            i += 1
            continue

        # Table detection (header row followed by delimiter row)
        if "|" in stripped and i + 1 < len(lines) and re.match(r"^\s*\|?\s*[-:]+[-| :]*$", lines[i + 1].strip()):
            table_lines = []
            while i < len(lines) and "|" in lines[i].strip():
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in table_lines:
                if re.match(r"^\s*\|?\s*[-:]+[-| :]*$", tl):
                    continue
                cells = [c.strip() for c in tl.strip("|").split("|")]
                rows.append(cells)
            blocks.append({"type": "table", "rows": rows})
            continue

        # Bullet list
        if stripped.startswith(("- ", "* ", "+ ")):
            items = []
            while i < len(lines) and lines[i].strip().startswith(("- ", "* ", "+ ")):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            blocks.append({"type": "list", "items": items, "ordered": False})
            continue

        # Numbered list
        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "list", "items": items, "ordered": True})
            continue

        # Blockquote
        if stripped.startswith(">"):
            quotes = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quotes.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            blocks.append({"type": "blockquote", "text": " ".join(quotes)})
            continue

        # Standard paragraph (accumulate non-empty, non-block lines)
        para = [stripped]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].strip().startswith(("#", "-", "*", "+", ">", "|"))
            and not re.match(r"^\d+\.\s+", lines[i].strip())
        ):
            para.append(lines[i].strip())
            i += 1
        blocks.append({"type": "paragraph", "text": " ".join(para)})

    return blocks


def parse_report_markdown(markdown_text: str, default_title: str = "Research & Operations Report") -> ReportData:
    """Extract structured report data from markdown, preserving backward compatibility."""
    lines = markdown_text.splitlines()

    title = default_title
    for line in lines:
        if line.startswith("# "):
            title = line.lstrip("# ").strip()
            break

    period = datetime.now().strftime("%B %Y")
    entity = "AgenticX Intelligence Platform"
    classification = "Official Research & Analysis Report"

    for line in lines:
        if "Reporting Period" in line or "Period" in line:
            m = re.search(r"(?:Reporting Period|Period)\*\*?:\s*([^–\n\)]+)", line)
            if m:
                period = m.group(1).strip()
        elif "Entity" in line:
            m = re.search(r"Entity\*\*?:\s*([^–\n\)]+)", line)
            if m:
                entity = m.group(1).strip()
        elif "Classification" in line:
            m = re.search(r"Classification\*\*?:\s*(.+)", line)
            if m:
                classification = m.group(1).strip()

    blocks = parse_markdown_blocks(markdown_text)
    kpis: list[dict[str, str]] = []
    risks: list[dict[str, str]] = []
    recommendations: list[str] = []

    for b in blocks:
        if b["type"] == "table" and b.get("rows"):
            rows = b["rows"]
            if len(rows) > 1:
                headers = rows[0]
                for r in rows[1:]:
                    metric_name = r[0] if len(r) > 0 else "Metric"
                    metric_val = r[1] if len(r) > 1 else ""
                    metric_chg = r[2] if len(r) > 2 else ""
                    metric_note = r[3] if len(r) > 3 else ""
                    kpis.append({
                        "metric": metric_name,
                        "value": metric_val,
                        "change": metric_chg,
                        "status": "neutral",
                        "note": metric_note or f"Spec: {', '.join(r[1:])}",
                    })
        elif b["type"] == "list":
            for item in b.get("items", []):
                clean_item = clean_inline_md(item)
                if any(w in clean_item.lower() for w in ("recommend", "action", "should", "suggest", "best")):
                    recommendations.append(clean_item)
                elif any(w in clean_item.lower() for w in ("risk", "vulnerability", "covenant", "critical", "warning")):
                    risks.append({
                        "id": f"R-0{len(risks)+1}",
                        "title": clean_item[:50],
                        "severity": "HIGH",
                        "impact": clean_item,
                        "mitigation": "Review operational constraints and guidelines.",
                    })

    if not recommendations:
        recommendations = ["Review detailed findings and implement findings according to specifications."]

    return ReportData(
        title=title,
        subtitle="Comprehensive Research, Specifications & Analytical Findings",
        period=period,
        entity=entity,
        classification=classification,
        kpis=kpis,
        covenants=[],
        concentration=[],
        risks=risks,
        recommendations=recommendations,
    )


# =============================================================================
# 2. ANIMATED PRESENTATION DECK (.HTML) GENERATOR
# =============================================================================

def generate_animated_html_deck(markdown_content: str, title: str = "") -> str:
    """Generate an interactive, high-contrast animated presentation deck matching research content."""
    blocks = parse_markdown_blocks(markdown_content)

    doc_title = title
    if not doc_title:
        for b in blocks:
            if b["type"] == "heading" and b["level"] == 1:
                doc_title = b["text"]
                break
    doc_title = doc_title or "Executive Presentation"

    # Split blocks into slide sections based on H1 or H2
    slides: list[dict[str, Any]] = []
    current_slide: dict[str, Any] = {"title": doc_title, "blocks": []}

    for b in blocks:
        if b["type"] == "heading" and b["level"] <= 2:
            if current_slide["blocks"]:
                slides.append(current_slide)
            current_slide = {"title": b["text"], "blocks": []}
        else:
            current_slide["blocks"].append(b)

    if current_slide["blocks"] or not slides:
        slides.append(current_slide)

    slide_html_cards = []
    for idx, slide in enumerate(slides):
        is_first = idx == 0
        slide_title = html.escape(slide["title"])

        body_parts = []
        for b in slide["blocks"]:
            if b["type"] == "heading":
                h_tag = f"h{min(b['level'] + 1, 4)}"
                body_parts.append(f"<{h_tag} style='color:#38bdf8; margin: 12px 0 6px 0;'>{html.escape(b['text'])}</{h_tag}>")
            elif b["type"] == "paragraph":
                body_parts.append(f"<p style='color:#cbd5e1; line-height: 1.6; margin-bottom: 12px; font-size: 1.05rem;'>{html.escape(b['text'])}</p>")
            elif b["type"] == "list":
                list_tag = "ol" if b.get("ordered") else "ul"
                items_html = "".join(f"<li style='margin-bottom: 8px; color: #e2e8f0;'>{html.escape(item)}</li>" for item in b.get("items", []))
                body_parts.append(f"<{list_tag} style='margin-left: 24px; margin-bottom: 16px; font-size: 1rem;'>{items_html}</{list_tag}>")
            elif b["type"] == "table":
                rows = b.get("rows", [])
                if rows:
                    th_html = "".join(f"<th style='background:#0f172a; padding:10px 14px; text-align:left; border:1px solid #334155; color:#38bdf8;'>{html.escape(c)}</th>" for c in rows[0])
                    tr_html = ""
                    for r in rows[1:]:
                        tds = "".join(f"<td style='padding:8px 12px; border:1px solid #1e293b; color:#f1f5f9;'>{html.escape(c)}</td>" for c in r)
                        tr_html += f"<tr>{tds}</tr>"
                    body_parts.append(f"<div style='overflow-x:auto; margin:16px 0;'><table style='width:100%; border-collapse:collapse; font-size:0.9rem;'><thead><tr>{th_html}</tr></thead><tbody>{tr_html}</tbody></table></div>")
            elif b["type"] == "blockquote":
                body_parts.append(f"<blockquote style='border-left:4px solid #38bdf8; background:rgba(15,23,42,0.6); padding:12px 18px; margin:14px 0; border-radius:4px; color:#93c5fd;'>{html.escape(b['text'])}</blockquote>")

        content_html = "\n".join(body_parts) if body_parts else "<p style='color:#94a3b8; font-size: 1.1rem;'>Official Research Brief & Synthesis</p>"

        active_cls = "active" if is_first else ""
        slide_html_cards.append(f"""
    <section class="slide {active_cls}" data-slide="{idx}">
      <div class="slide-card">
        <div class="badge-tag">Section {idx + 1} of {len(slides)}</div>
        <h2 class="slide-header">{slide_title}</h2>
        <div class="slide-body">
          {content_html}
        </div>
      </div>
    </section>""")

    rendered_slides = "\n".join(slide_html_cards)

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{html.escape(doc_title)} — Interactive Deck</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@400;600;700;800&family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
  <style>
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{
      background: #090d16;
      color: #f8fafc;
      font-family: 'Inter', sans-serif;
      min-height: 100vh;
      display: flex;
      flex-direction: column;
      overflow: hidden;
    }}
    header {{
      padding: 18px 32px;
      display: flex;
      justify-content: space-between;
      align-items: center;
      border-bottom: 1px solid rgba(255, 255, 255, 0.08);
      background: rgba(15, 23, 42, 0.85);
      backdrop-filter: blur(12px);
      z-index: 10;
    }}
    .brand {{
      font-family: 'Outfit', sans-serif;
      font-weight: 800;
      font-size: 1.25rem;
      letter-spacing: -0.02em;
      color: #38bdf8;
    }}
    .progress-pill {{
      font-size: 0.85rem;
      font-weight: 600;
      background: rgba(56, 189, 248, 0.1);
      color: #38bdf8;
      border: 1px solid rgba(56, 189, 248, 0.3);
      padding: 4px 14px;
      border-radius: 9999px;
    }}
    main {{
      flex: 1;
      position: relative;
      display: flex;
      align-items: center;
      justify-content: center;
      padding: 24px;
    }}
    .slide {{
      position: absolute;
      inset: 24px;
      display: flex;
      align-items: center;
      justify-content: center;
      opacity: 0;
      pointer-events: none;
      transform: translateY(20px) scale(0.98);
      transition: all 0.35s cubic-bezier(0.16, 1, 0.3, 1);
    }}
    .slide.active {{
      opacity: 1;
      pointer-events: auto;
      transform: translateY(0) scale(1);
    }}
    .slide-card {{
      background: rgba(15, 23, 42, 0.75);
      border: 1px solid rgba(255, 255, 255, 0.1);
      box-shadow: 0 25px 50px -12px rgba(0, 0, 0, 0.6);
      border-radius: 16px;
      padding: 40px;
      max-width: 960px;
      width: 100%;
      max-height: 82vh;
      overflow-y: auto;
      backdrop-filter: blur(16px);
    }}
    .badge-tag {{
      display: inline-block;
      font-size: 0.75rem;
      font-weight: 700;
      color: #38bdf8;
      text-transform: uppercase;
      letter-spacing: 0.08em;
      margin-bottom: 8px;
    }}
    .slide-header {{
      font-family: 'Outfit', sans-serif;
      font-size: 2rem;
      font-weight: 800;
      color: #ffffff;
      margin-bottom: 20px;
      line-height: 1.2;
      border-bottom: 1px solid rgba(255, 255, 255, 0.08);
      padding-bottom: 12px;
    }}
    .slide-body {{
      font-size: 1.05rem;
    }}
    footer {{
      padding: 16px 32px;
      display: flex;
      justify-content: space-between;
      align-items: center;
      border-top: 1px solid rgba(255, 255, 255, 0.08);
      background: rgba(15, 23, 42, 0.85);
      z-index: 10;
    }}
    .nav-btn {{
      background: #1e293b;
      color: #f8fafc;
      border: 1px solid rgba(255, 255, 255, 0.15);
      padding: 8px 18px;
      border-radius: 8px;
      cursor: pointer;
      font-weight: 600;
      font-size: 0.9rem;
      transition: all 0.2s ease;
    }}
    .nav-btn:hover:not(:disabled) {{
      background: #38bdf8;
      color: #0f172a;
      border-color: #38bdf8;
    }}
    .nav-btn:disabled {{
      opacity: 0.4;
      cursor: not-allowed;
    }}
    .key-hints {{
      color: #64748b;
      font-size: 0.8rem;
    }}
  </style>
</head>
<body>
  <header>
    <div class="brand">AgenticX Presentation</div>
    <div class="progress-pill" id="counter">Slide 1 of {len(slides)}</div>
  </header>

  <main>
    {rendered_slides}
  </main>

  <footer>
    <div class="key-hints">Tip: Use Arrow Keys (← / →) or Spacebar</div>
    <div style="display: flex; gap: 10px;">
      <button class="nav-btn" id="prevBtn" onclick="changeSlide(-1)" disabled>Previous</button>
      <button class="nav-btn" id="nextBtn" onclick="changeSlide(1)">Next</button>
    </div>
  </footer>

  <script>
    let current = 0;
    const total = {len(slides)};
    const slides = document.querySelectorAll('.slide');
    const counter = document.getElementById('counter');
    const prevBtn = document.getElementById('prevBtn');
    const nextBtn = document.getElementById('nextBtn');

    function updateSlide() {{
      slides.forEach((s, idx) => {{
        s.classList.toggle('active', idx === current);
      }});
      counter.textContent = `Slide ${{current + 1}} of ${{total}}`;
      prevBtn.disabled = current === 0;
      nextBtn.disabled = current === total - 1;
    }}

    function changeSlide(direction) {{
      const next = current + direction;
      if (next >= 0 && next < total) {{
        current = next;
        updateSlide();
      }}
    }}

    document.addEventListener('keydown', (e) => {{
      if (e.key === 'ArrowRight' || e.key === ' ') {{
        changeSlide(1);
      }} else if (e.key === 'ArrowLeft') {{
        changeSlide(-1);
      }}
    }});
  </script>
</body>
</html>"""


# =============================================================================
# 3. EXECUTIVE PRESENTATION (.PPTX) GENERATOR
# =============================================================================

def generate_pptx(markdown_content: str, title: str = "", output_path: Path | None = None) -> Path:
    """Generate an executive, high-contrast 16:9 PowerPoint deck from markdown content."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    if output_path is None:
        output_path = settings.workspace_path / "presentation.pptx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    blocks = parse_markdown_blocks(markdown_content)

    doc_title = title
    if not doc_title:
        for b in blocks:
            if b["type"] == "heading" and b["level"] == 1:
                doc_title = b["text"]
                break
    doc_title = doc_title or "Executive Presentation"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    COLOR_BG = RGBColor(11, 19, 43)          # Deep Navy
    COLOR_CARD = RGBColor(20, 32, 58)        # Card Navy
    COLOR_CYAN = RGBColor(0, 240, 255)       # Electric Cyan
    COLOR_WHITE = RGBColor(255, 255, 255)    # Pure White
    COLOR_MUTED = RGBColor(148, 163, 184)    # Slate Muted

    def apply_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    # Slide 1: Title Slide
    slide1 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(slide1)

    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "OFFICIAL REPORT • AGENTICX RESEARCH ENGINE"
    p_badge.font.name = "Calibri"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_CYAN

    p_main = tf1.add_paragraph()
    p_main.text = clean_inline_md(doc_title)
    p_main.font.name = "Calibri"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_WHITE
    p_main.space_before = Pt(12)

    p_sub = tf1.add_paragraph()
    p_sub.text = f"Compiled autonomously | {datetime.now().strftime('%B %d, %Y')}"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = COLOR_MUTED
    p_sub.space_before = Pt(12)

    # Group into subsequent slides
    section_slides: list[dict[str, Any]] = []
    curr: dict[str, Any] = {"title": "Summary & Findings", "blocks": []}

    for b in blocks:
        if b["type"] == "heading" and b["level"] <= 2:
            if curr["blocks"]:
                section_slides.append(curr)
            curr = {"title": clean_inline_md(b["text"]), "blocks": []}
        else:
            curr["blocks"].append(b)

    if curr["blocks"]:
        section_slides.append(curr)

    for s_data in section_slides:
        slide = prs.slides.add_slide(blank_layout)
        apply_slide_bg(slide)

        # Header Box
        h_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        htf = h_box.text_frame
        htf.word_wrap = True
        hp_cat = htf.paragraphs[0]
        hp_cat.text = "AGENTICX RESEARCH • DETAILED FINDINGS"
        hp_cat.font.name = "Calibri"
        hp_cat.font.size = Pt(10)
        hp_cat.font.bold = True
        hp_cat.font.color.rgb = COLOR_CYAN

        hp_title = htf.add_paragraph()
        hp_title.text = s_data["title"]
        hp_title.font.name = "Calibri"
        hp_title.font.size = Pt(22)
        hp_title.font.bold = True
        hp_title.font.color.rgb = COLOR_WHITE
        hp_title.space_before = Pt(4)

        # Content Card
        c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = COLOR_CARD
        c_shape.line.color.rgb = RGBColor(42, 60, 95)
        c_shape.line.width = Pt(1)

        # Content placement
        # Check if table block exists
        table_block = next((b for b in s_data["blocks"] if b["type"] == "table"), None)
        if table_block and table_block.get("rows"):
            rows = table_block["rows"][:8]  # Limit to 8 rows for slide readability
            num_rows = len(rows)
            num_cols = max(len(r) for r in rows)
            table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(1.1), Inches(1.8), Inches(11.133), Inches(4.7))
            tbl = table_shape.table

            for r_idx, row in enumerate(rows):
                for c_idx in range(num_cols):
                    cell = tbl.cell(r_idx, c_idx)
                    cell.text = row[c_idx] if c_idx < len(row) else ""
                    for p in cell.text_frame.paragraphs:
                        p.font.name = "Calibri"
                        p.font.size = Pt(10 if r_idx == 0 else 9)
                        if r_idx == 0:
                            p.font.bold = True
                            p.font.color.rgb = COLOR_CYAN
                        else:
                            p.font.color.rgb = COLOR_WHITE
        else:
            tf_body = c_shape.text_frame
            tf_body.word_wrap = True
            tf_body.margin_left = Inches(0.3)
            tf_body.margin_top = Inches(0.3)
            tf_body.margin_right = Inches(0.3)

            first_p = True
            for b in s_data["blocks"]:
                if b["type"] == "heading":
                    p = tf_body.paragraphs[0] if first_p else tf_body.add_paragraph()
                    first_p = False
                    p.text = clean_inline_md(b["text"])
                    p.font.name = "Calibri"
                    p.font.size = Pt(15)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_CYAN
                    p.space_before = Pt(8)
                elif b["type"] in ("paragraph", "blockquote"):
                    p = tf_body.paragraphs[0] if first_p else tf_body.add_paragraph()
                    first_p = False
                    p.text = clean_inline_md(b["text"])
                    p.font.name = "Calibri"
                    p.font.size = Pt(12)
                    p.font.color.rgb = COLOR_WHITE
                    p.space_before = Pt(6)
                elif b["type"] == "list":
                    for item in b.get("items", []):
                        p = tf_body.paragraphs[0] if first_p else tf_body.add_paragraph()
                        first_p = False
                        p.text = "• " + clean_inline_md(item)
                        p.font.name = "Calibri"
                        p.font.size = Pt(11.5)
                        p.font.color.rgb = RGBColor(226, 232, 240)
                        p.space_before = Pt(4)

    prs.save(str(output_path))
    logger.info("Generated presentation (.pptx): %s", output_path)
    return output_path


# =============================================================================
# 4. EXECUTIVE WORD DOCUMENT (.DOCX) GENERATOR
# =============================================================================

def generate_docx(markdown_content: str, title: str = "", output_path: Path | None = None) -> Path:
    """Generate a clean, publication-ready Word Document (.docx) matching the markdown content."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Inches, Pt, RGBColor

    if output_path is None:
        output_path = settings.workspace_path / "report.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    blocks = parse_markdown_blocks(markdown_content)

    doc_title = title
    if not doc_title:
        for b in blocks:
            if b["type"] == "heading" and b["level"] == 1:
                doc_title = b["text"]
                break
    doc_title = doc_title or "Executive Research Report"

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"AgenticX Research Report — {doc_title}")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("CONFIDENTIAL • AGENTICX AUTONOMOUS RESEARCH")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Document Title Block
    p_tag = doc.add_paragraph()
    r_tag = p_tag.add_run("OFFICIAL RESEARCH REPORT • AGENTICX")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(9)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(37, 99, 235)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run(clean_inline_md(doc_title))
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)

    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y')}")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    p_meta.paragraph_format.space_after = Pt(14)

    # Divider bar
    p_div = doc.add_paragraph()
    p_div_run = p_div.add_run("―" * 58)
    p_div_run.font.color.rgb = RGBColor(203, 213, 225)
    p_div.paragraph_format.space_after = Pt(14)

    def set_cell_background(cell, fill_hex: str):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Render Markdown blocks
    for b in blocks:
        if b["type"] == "heading":
            level = b["level"]
            text = clean_inline_md(b["text"])
            if level == 1 and text == clean_inline_md(doc_title):
                continue  # Avoid repeating main title

            h = doc.add_heading(level=min(level, 3))
            rh = h.add_run(text)
            rh.font.name = "Calibri"
            rh.font.bold = True
            if level == 2:
                rh.font.size = Pt(14)
                rh.font.color.rgb = RGBColor(15, 23, 42)
                h.paragraph_format.space_before = Pt(14)
                h.paragraph_format.space_after = Pt(6)
            else:
                rh.font.size = Pt(12)
                rh.font.color.rgb = RGBColor(30, 41, 59)
                h.paragraph_format.space_before = Pt(10)
                h.paragraph_format.space_after = Pt(4)

        elif b["type"] == "paragraph":
            p = doc.add_paragraph(clean_inline_md(b["text"]))
            p.style.font.name = "Calibri"
            p.style.font.size = Pt(10.5)
            p.paragraph_format.space_after = Pt(8)

        elif b["type"] == "blockquote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(clean_inline_md(b["text"]))
            r.font.name = "Calibri"
            r.font.italic = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(2, 132, 199)

        elif b["type"] == "list":
            style_name = "List Number" if b.get("ordered") else "List Bullet"
            for item in b.get("items", []):
                p = doc.add_paragraph(style=style_name)
                r = p.add_run(clean_inline_md(item))
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                p.paragraph_format.space_after = Pt(3)

        elif b["type"] == "table":
            rows = b.get("rows", [])
            if rows:
                num_cols = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True

                for r_idx, r_data in enumerate(rows):
                    row_cells = tbl.rows[r_idx].cells
                    for c_idx in range(num_cols):
                        cell_text = clean_inline_md(r_data[c_idx]) if c_idx < len(r_data) else ""
                        row_cells[c_idx].text = cell_text
                        cp = row_cells[c_idx].paragraphs[0]
                        cp.paragraph_format.space_after = Pt(2)
                        cp.paragraph_format.space_before = Pt(2)
                        if cp.runs:
                            crun = cp.runs[0]
                            crun.font.name = "Calibri"
                            crun.font.size = Pt(9)
                            if r_idx == 0:
                                crun.font.bold = True
                                crun.font.color.rgb = RGBColor(255, 255, 255)
                                set_cell_background(row_cells[c_idx], "0F172A")
                            else:
                                crun.font.color.rgb = RGBColor(30, 41, 59)
                                if r_idx % 2 == 1:
                                    set_cell_background(row_cells[c_idx], "F8FAFC")

                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_after = Pt(10)

    doc.save(str(output_path))
    logger.info("Generated Word document (.docx): %s", output_path)
    return output_path


# =============================================================================
# 5. PUBLICATION-QUALITY EXECUTIVE PDF (.PDF) GENERATOR
# =============================================================================

def generate_pdf(markdown_content: str, title: str = "", output_path: Path | None = None) -> Path:
    """Generate a vector-sharp, publication-quality executive PDF using Chrome print or ReportLab."""
    from markdown_it import MarkdownIt

    if output_path is None:
        output_path = settings.workspace_path / "report.pdf"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Extract display title and body
    lines = markdown_content.splitlines()
    display_title = title
    body_lines: list[str] = []
    found_h1 = False

    for line in lines:
        if not found_h1 and line.strip().startswith("# "):
            if not display_title:
                display_title = line.strip().lstrip("# ").strip()
            found_h1 = True
        else:
            body_lines.append(line)

    display_title = display_title or "Executive Research Report"
    body_md = "\n".join(body_lines)

    md = MarkdownIt("commonmark").enable("table").enable("strikethrough")
    body_html = md.render(body_md)
    date_str = datetime.now().strftime("%B %d, %Y")

    # Generate dedicated printable HTML
    printable_html = f"""<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8">
  <title>{html.escape(display_title)}</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@600;700;800&family=Inter:wght@400;500;600;700&family=JetBrains+Mono:wght@400;500;600&display=swap" rel="stylesheet">
  <style>
    @page {{
      size: A4 portrait;
      margin: 18mm 16mm 18mm 16mm;
      @bottom-center {{
        content: "Page " counter(page) " of " counter(pages);
        font-family: 'Inter', sans-serif;
        font-size: 8pt;
        color: #94a3b8;
      }}
    }}
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{
      font-family: 'Inter', sans-serif;
      color: #0f172a;
      line-height: 1.55;
      font-size: 9.5pt;
      background: #ffffff;
      padding: 0;
    }}
    .header-banner {{
      border-bottom: 2px solid #0f172a;
      padding-bottom: 12px;
      margin-bottom: 20px;
    }}
    .class-tag {{
      display: inline-block;
      font-size: 7.5pt;
      font-weight: 700;
      color: #2563eb;
      text-transform: uppercase;
      letter-spacing: 0.08em;
      margin-bottom: 6px;
    }}
    h1 {{
      font-family: 'Outfit', sans-serif;
      font-size: 22pt;
      font-weight: 800;
      color: #0f172a;
      line-height: 1.15;
      margin-bottom: 6px;
    }}
    .meta-line {{
      font-size: 8.5pt;
      color: #64748b;
    }}
    h2 {{
      font-family: 'Outfit', sans-serif;
      font-size: 13pt;
      font-weight: 700;
      color: #0f172a;
      margin-top: 20px;
      margin-bottom: 8px;
      border-bottom: 1px solid #e2e8f0;
      padding-bottom: 4px;
      page-break-after: avoid;
    }}
    h3 {{
      font-family: 'Outfit', sans-serif;
      font-size: 11pt;
      font-weight: 600;
      color: #1e293b;
      margin-top: 14px;
      margin-bottom: 4px;
      page-break-after: avoid;
    }}
    p, ul, ol {{
      margin-bottom: 10px;
      color: #334155;
    }}
    ul, ol {{
      margin-left: 20px;
    }}
    li {{
      margin-bottom: 4px;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin: 14px 0 16px 0;
      font-size: 8.5pt;
      page-break-inside: avoid;
    }}
    th {{
      background: #0f172a;
      color: #ffffff;
      font-weight: 600;
      text-align: left;
      padding: 8px 10px;
      border: 1px solid #0f172a;
    }}
    td {{
      padding: 8px 10px;
      border: 1px solid #e2e8f0;
      color: #1e293b;
      vertical-align: top;
    }}
    tr:nth-child(even) td {{
      background: #f8fafc;
    }}
    blockquote {{
      background: #f0f9ff;
      border-left: 4px solid #0284c7;
      padding: 10px 14px;
      border-radius: 4px;
      margin: 14px 0;
      font-size: 9pt;
      color: #0369a1;
    }}
    code {{
      font-family: 'JetBrains Mono', monospace;
      font-size: 8.5pt;
      background: #f1f5f9;
      padding: 2px 4px;
      border-radius: 3px;
      color: #0f172a;
    }}
    pre {{
      background: #0f172a;
      color: #f8fafc;
      padding: 12px 14px;
      border-radius: 6px;
      overflow-x: auto;
      margin: 12px 0;
    }}
    pre code {{
      background: none;
      padding: 0;
      color: inherit;
    }}
    a {{
      color: #2563eb;
      text-decoration: none;
      font-weight: 500;
    }}
  </style>
</head>
<body>
  <div class="header-banner">
    <div class="class-tag">Official Research Report • AgenticX Engine</div>
    <h1>{html.escape(display_title)}</h1>
    <div class="meta-line">Generated by AgenticX Intelligence Platform | {date_str}</div>
  </div>
  {body_html}
</body>
</html>"""

    # 1. Print to PDF using Headless Chrome
    chrome_binary = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    if Path(chrome_binary).exists():
        with tempfile.NamedTemporaryFile(suffix=".html", mode="w", encoding="utf-8", delete=False) as tf:
            tf.write(printable_html)
            temp_html_path = tf.name

        try:
            cmd = [
                chrome_binary,
                "--headless",
                "--disable-gpu",
                f"--print-to-pdf={output_path}",
                "--no-pdf-header-footer",
                f"file://{temp_html_path}",
            ]
            subprocess.run(cmd, check=True, timeout=30, capture_output=True)
            logger.info("Generated PDF via headless Chrome: %s", output_path)
            return output_path
        except Exception as exc:
            logger.warning("Headless Chrome PDF export failed (%s), trying ReportLab fallback", exc)
        finally:
            try:
                os.remove(temp_html_path)
            except OSError:
                pass

    # 2. ReportLab Platypus Dynamic Fallback
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    doc = SimpleDocTemplate(str(output_path), pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("TitleStyle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor("#0F172A"))
    h2_style = ParagraphStyle("H2Style", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=17, textColor=colors.HexColor("#0F172A"), spaceBefore=12, spaceAfter=6)
    h3_style = ParagraphStyle("H3Style", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=11, leading=15, textColor=colors.HexColor("#1E293B"), spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle("BodyStyle", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=14, textColor=colors.HexColor("#334155"), spaceAfter=6)
    bullet_style = ParagraphStyle("BulletStyle", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor("#334155"), leftIndent=15, spaceAfter=3)

    blocks = parse_markdown_blocks(markdown_content)
    story = []

    story.append(Paragraph(clean_inline_md(display_title), title_style))
    story.append(Spacer(1, 6))
    meta_style = ParagraphStyle("MetaStyle", parent=styles["Normal"], fontName="Helvetica", fontSize=8.5, textColor=colors.HexColor("#64748B"))
    story.append(Paragraph(f"Generated by AgenticX Engine | {date_str}", meta_style))
    story.append(Spacer(1, 14))

    for b in blocks:
        if b["type"] == "heading":
            htext = clean_inline_md(b["text"])
            if htext == clean_inline_md(display_title):
                continue
            if b["level"] <= 2:
                story.append(Paragraph(htext, h2_style))
            else:
                story.append(Paragraph(htext, h3_style))
        elif b["type"] in ("paragraph", "blockquote"):
            story.append(Paragraph(clean_inline_md(b["text"]), body_style))
        elif b["type"] == "list":
            for item in b.get("items", []):
                story.append(Paragraph("• " + clean_inline_md(item), bullet_style))
        elif b["type"] == "table":
            rows = b.get("rows", [])
            if rows:
                table_data = []
                for r in rows:
                    table_data.append([Paragraph(clean_inline_md(c), body_style) for c in r])
                t = Table(table_data)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F8FAFC")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ]))
                story.append(t)
                story.append(Spacer(1, 10))

    doc.build(story)
    logger.info("Generated PDF via ReportLab: %s", output_path)
    return output_path
